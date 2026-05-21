"""NBA 2026 ECF Game 2 ニックス対キャバリアーズ - WordおよびExcelファイル生成"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT_DIR = "/home/user/search-launcher"


def make_word():
    doc = Document()

    # Title
    title = doc.add_heading("2026 NBAイースタン・カンファレンスファイナル Game 2", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("ニューヨーク・ニックス 対 クリーブランド・キャバリアーズ")
    run.bold = True
    run.font.size = Pt(14)

    # 1. 試合概要
    doc.add_heading("1. 試合概要", level=1)
    info = doc.add_table(rows=0, cols=2)
    info.style = "Light Grid Accent 1"
    rows = [
        ("試合日時（米国東部時間）", "2026年5月21日（木）20:00 ET"),
        ("試合日時（日本時間）", "2026年5月22日（金）9:00 JST"),
        ("会場", "マディソン・スクエア・ガーデン（ニューヨーク）"),
        ("シリーズ", "イースタン・カンファレンスファイナル Game 2"),
        ("シリーズ状況", "ニックスが1勝0敗でリード"),
        ("中継（米国）", "ESPN"),
    ]
    for k, v in rows:
        row = info.add_row().cells
        row[0].text = k
        row[1].text = v

    # 2. シリーズ日程
    doc.add_heading("2. シリーズ日程（米国東部時間）", level=1)
    sched = doc.add_table(rows=1, cols=3)
    sched.style = "Light Grid Accent 1"
    hdr = sched.rows[0].cells
    hdr[0].text = "Game"
    hdr[1].text = "日付（ET）"
    hdr[2].text = "中継"
    schedule_rows = [
        ("Game 1", "2026年5月19日（火）", "ESPN"),
        ("Game 2", "2026年5月21日（木）20:00", "ESPN"),
        ("Game 3", "2026年5月23日（土）20:00", "ABC"),
        ("Game 4", "2026年5月25日（月）20:00", "ESPN"),
        ("Game 5（必要時）", "2026年5月27日（水）", "ESPN"),
        ("Game 6（必要時）", "2026年5月29日（金）", "ESPN"),
        ("Game 7（必要時）", "2026年5月31日（日）", "ESPN"),
    ]
    for g, d, n in schedule_rows:
        r = sched.add_row().cells
        r[0].text = g
        r[1].text = d
        r[2].text = n

    # 3. Game 1 結果（事実のみ）
    doc.add_heading("3. Game 1 結果（2026年5月19日、MSG）", level=1)
    p = doc.add_paragraph()
    p.add_run("最終スコア：").bold = True
    p.add_run("ニックス 115 – キャバリアーズ 104（OT／延長）")

    p2 = doc.add_paragraph()
    p2.add_run("試合経過の主なポイント：").bold = True
    bullets = [
        "前半終了時：キャバリアーズが48-46でリード（キャバリアーズが前半終盤に18-6のラン）。",
        "第4Q残り7:52時点：キャバリアーズが93-71と22点リード。",
        "そこからニックスが44-11のランで逆転、延長戦に持ち込み勝利。",
        "ニックスはOT中14得点のうち9点をOG・アヌノビーが獲得。",
        "NBA公式によると、この4Qの22点差からの逆転は、プレーバイプレー記録が残る1997年以降のカンファレンスファイナルにおける第4クォーター最大の逆転劇。",
        "Game 1終盤、キャバリアーズはニックスの逆転中に約22%のFG成功率に終わった。",
    ]
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")

    # 4. Game 1 個人スタッツ（確認済み）
    doc.add_heading("4. Game 1 主要プレーヤー成績（確認済みの数値のみ）", level=1)

    doc.add_paragraph().add_run("ニューヨーク・ニックス").bold = True
    knicks_tbl = doc.add_table(rows=1, cols=2)
    knicks_tbl.style = "Light Grid Accent 1"
    knicks_tbl.rows[0].cells[0].text = "選手"
    knicks_tbl.rows[0].cells[1].text = "成績"
    knicks_stats = [
        ("ジェイレン・ブランソン", "38得点、6アシスト、5リバウンド"),
        ("ミカル・ブリッジス", "18得点"),
        ("カール=アンソニー・タウンズ", "13得点、13リバウンド、5アシスト、1ブロック（ダブルダブル）"),
        ("OG・アヌノビー", "13得点（うちOTで9点）、OTで3リバウンド"),
    ]
    for n, s in knicks_stats:
        r = knicks_tbl.add_row().cells
        r[0].text = n
        r[1].text = s

    doc.add_paragraph().add_run("クリーブランド・キャバリアーズ").bold = True
    cavs_tbl = doc.add_table(rows=1, cols=2)
    cavs_tbl.style = "Light Grid Accent 1"
    cavs_tbl.rows[0].cells[0].text = "選手"
    cavs_tbl.rows[0].cells[1].text = "成績"
    cavs_stats = [
        ("ドノバン・ミッチェル", "29得点（FG 12-23、3P 4-11、FT 1-1）、6スティール（自己最多）、5リバウンド、3アシスト、1ブロック、出場41分"),
        ("エヴァン・モブリー", "15得点、14リバウンド"),
        ("ジェームズ・ハーデン", "9得点"),
        ("ジャレット・アレン", "8得点"),
    ]
    for n, s in cavs_stats:
        r = cavs_tbl.add_row().cells
        r[0].text = n
        r[1].text = s

    # 5. Game 2 インジュリーレポート
    doc.add_heading("5. Game 2 インジュリーレポート（試合前情報）", level=1)
    inj = doc.add_table(rows=1, cols=3)
    inj.style = "Light Grid Accent 1"
    inj.rows[0].cells[0].text = "選手"
    inj.rows[0].cells[1].text = "チーム"
    inj.rows[0].cells[2].text = "状況"
    inj_rows = [
        ("OG・アヌノビー", "ニックス", "出場可能。ただし5月6日のハムストリング肉離れから完全回復には至っていない"),
        ("ラリー・ナンス・ジュニア", "キャバリアーズ", "questionable（出場微妙）"),
    ]
    for n, t, s in inj_rows:
        r = inj.add_row().cells
        r[0].text = n
        r[1].text = t
        r[2].text = s
    doc.add_paragraph("※両チームともそれ以外の主要選手は出場可能。")

    # 6. 出典
    doc.add_heading("6. 出典", level=1)
    sources = [
        "NBA.com公式 2026 East Final 特集ページ：https://www.nba.com/playoffs/2026/east-final",
        "NBA.com ニックス公式 Game 2 ゲームノート：https://www.nba.com/knicks/news/knicks-vs-cavaliers-game-notes-may-21-2026-eastern-conference-finals-game-2",
        "NBA.com Game 1 ライブアップデート／リキャップ：https://www.nba.com/news/live-updates-2026-nba-playoffs-eastern-conference-finals-cavaliers-knicks-in-the-garden",
        "NBA.com 「4 takeaways」 Game 1：https://www.nba.com/news/4-takeaways-cavaliers-knicks-game-1",
        "ESPN Game 1 ファイナルスコア／リキャップ：https://www.espn.com/nba/game/_/gameId/401873341/cavaliers-knicks",
        "ESPN 2026 NBAプレーオフ スケジュール：https://www.espn.com/nba/story/_/id/48419498/nba-playoffs-2026-play-finals-schedule-scores-news-highlights-bracket-dates",
        "Yahoo Sports Game 1 リキャップ：https://sports.yahoo.com/articles/knicks-vs-cavaliers-live-score-235147883.html",
        "amNewYork Game 1 リキャップ：https://www.amny.com/sports/knicks-cavaliers-game1-ecf-5-19-26/",
    ]
    for s in sources:
        doc.add_paragraph(s, style="List Bullet")

    # 注記
    doc.add_heading("注記", level=1)
    note = doc.add_paragraph()
    note.add_run(
        "本ドキュメントはGame 2開始前時点の情報をまとめたもので、確認済みの事実のみを記載しています。"
        "予想・予測・推測は一切含めていません。試合直前の追加情報（スターティングメンバー、最終インジュリー判断等）は試合直前に各公式ソースで再確認してください。"
    )

    path = f"{OUT_DIR}/NBA_ECF_Game2_Knicks_vs_Cavaliers.docx"
    doc.save(path)
    return path


def make_excel():
    wb = Workbook()

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    knicks_fill = PatternFill(start_color="F58426", end_color="F58426", fill_type="solid")
    cavs_fill = PatternFill(start_color="860038", end_color="860038", fill_type="solid")
    team_font = Font(bold=True, color="FFFFFF", size=12)
    thin = Side(border_style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def apply_header(ws, row, cols):
        for c in range(1, cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border

    def apply_body_border(ws, r1, r2, c1, c2):
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(row=r, column=c).border = border
                if ws.cell(row=r, column=c).alignment.horizontal is None:
                    ws.cell(row=r, column=c).alignment = left

    # Sheet 1: 試合概要
    ws1 = wb.active
    ws1.title = "試合概要"
    ws1["A1"] = "2026 NBAイースタン・カンファレンスファイナル Game 2"
    ws1["A1"].font = Font(bold=True, size=16)
    ws1.merge_cells("A1:B1")
    ws1["A2"] = "ニューヨーク・ニックス 対 クリーブランド・キャバリアーズ"
    ws1["A2"].font = Font(bold=True, size=12)
    ws1.merge_cells("A2:B2")

    ws1["A4"] = "項目"
    ws1["B4"] = "内容"
    apply_header(ws1, 4, 2)

    overview = [
        ("試合日時（米国東部時間）", "2026年5月21日（木）20:00 ET"),
        ("試合日時（日本時間）", "2026年5月22日（金）9:00 JST"),
        ("会場", "マディソン・スクエア・ガーデン（ニューヨーク）"),
        ("シリーズ", "イースタン・カンファレンスファイナル Game 2"),
        ("シリーズ状況", "ニックス 1勝 0敗"),
        ("中継（米国）", "ESPN"),
        ("Game 1 結果", "ニックス 115 – キャバリアーズ 104（OT、2026年5月19日、MSG）"),
    ]
    for i, (k, v) in enumerate(overview, start=5):
        ws1.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws1.cell(row=i, column=2, value=v)
    apply_body_border(ws1, 5, 4 + len(overview), 1, 2)
    ws1.column_dimensions["A"].width = 30
    ws1.column_dimensions["B"].width = 60

    # Sheet 2: シリーズ日程
    ws2 = wb.create_sheet("シリーズ日程")
    ws2["A1"] = "シリーズ日程（米国東部時間）"
    ws2["A1"].font = Font(bold=True, size=14)
    ws2.merge_cells("A1:C1")
    headers = ["Game", "日付（ET）", "中継"]
    for i, h in enumerate(headers, start=1):
        ws2.cell(row=3, column=i, value=h)
    apply_header(ws2, 3, 3)
    schedule = [
        ("Game 1", "2026年5月19日（火）", "ESPN"),
        ("Game 2", "2026年5月21日（木）20:00", "ESPN"),
        ("Game 3", "2026年5月23日（土）20:00", "ABC"),
        ("Game 4", "2026年5月25日（月）20:00", "ESPN"),
        ("Game 5（必要時）", "2026年5月27日（水）", "ESPN"),
        ("Game 6（必要時）", "2026年5月29日（金）", "ESPN"),
        ("Game 7（必要時）", "2026年5月31日（日）", "ESPN"),
    ]
    for i, row in enumerate(schedule, start=4):
        for j, val in enumerate(row, start=1):
            cell = ws2.cell(row=i, column=j, value=val)
            cell.alignment = center
    apply_body_border(ws2, 4, 3 + len(schedule), 1, 3)
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 28
    ws2.column_dimensions["C"].width = 12

    # Sheet 3: Game 1 ニックス スタッツ
    ws3 = wb.create_sheet("Game1_ニックス")
    ws3["A1"] = "Game 1 ニューヨーク・ニックス 主要選手成績"
    ws3["A1"].font = Font(bold=True, size=14)
    ws3.merge_cells("A1:H1")
    ws3["A2"] = "2026年5月19日 / 最終スコア: ニックス 115 - キャバリアーズ 104 (OT)"
    ws3.merge_cells("A2:H2")
    ws3["A3"] = "ニューヨーク・ニックス"
    ws3["A3"].fill = knicks_fill
    ws3["A3"].font = team_font
    ws3["A3"].alignment = center
    ws3.merge_cells("A3:H3")

    headers3 = ["選手", "得点", "リバウンド", "アシスト", "スティール", "ブロック", "出場分", "備考"]
    for i, h in enumerate(headers3, start=1):
        ws3.cell(row=4, column=i, value=h)
    apply_header(ws3, 4, len(headers3))

    knicks = [
        ("ジェイレン・ブランソン", 38, 5, 6, None, None, None, "4Q大逆転を牽引"),
        ("ミカル・ブリッジス", 18, None, None, None, None, None, ""),
        ("カール=アンソニー・タウンズ", 13, 13, 5, None, 1, None, "ダブルダブル"),
        ("OG・アヌノビー", 13, None, None, None, None, None, "OTで9得点・3リバウンド、ハムストリング負傷から復帰"),
    ]
    for i, row in enumerate(knicks, start=5):
        for j, val in enumerate(row, start=1):
            cell = ws3.cell(row=i, column=j, value=val if val is not None else "")
            if j == 1:
                cell.alignment = left
            else:
                cell.alignment = center
    apply_body_border(ws3, 5, 4 + len(knicks), 1, len(headers3))
    widths3 = [24, 8, 12, 10, 10, 8, 10, 50]
    for i, w in enumerate(widths3, start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3["A10"] = "※空欄は本ファイル作成時点で公開情報で確認できなかった項目（推測しない方針のため空欄）。"
    ws3["A10"].font = Font(italic=True, size=9, color="666666")
    ws3.merge_cells("A10:H10")

    # Sheet 4: Game 1 キャバリアーズ スタッツ
    ws4 = wb.create_sheet("Game1_キャバリアーズ")
    ws4["A1"] = "Game 1 クリーブランド・キャバリアーズ 主要選手成績"
    ws4["A1"].font = Font(bold=True, size=14)
    ws4.merge_cells("A1:J1")
    ws4["A2"] = "2026年5月19日 / 最終スコア: ニックス 115 - キャバリアーズ 104 (OT)"
    ws4.merge_cells("A2:J2")
    ws4["A3"] = "クリーブランド・キャバリアーズ"
    ws4["A3"].fill = cavs_fill
    ws4["A3"].font = team_font
    ws4["A3"].alignment = center
    ws4.merge_cells("A3:J3")

    headers4 = ["選手", "得点", "リバウンド", "アシスト", "スティール", "ブロック", "出場分", "FG", "3P", "FT"]
    for i, h in enumerate(headers4, start=1):
        ws4.cell(row=4, column=i, value=h)
    apply_header(ws4, 4, len(headers4))

    cavs = [
        ("ドノバン・ミッチェル", 29, 5, 3, 6, 1, 41, "12-23", "4-11", "1-1"),
        ("エヴァン・モブリー", 15, 14, None, None, None, None, "", "", ""),
        ("ジェームズ・ハーデン", 9, None, None, None, None, None, "", "", ""),
        ("ジャレット・アレン", 8, None, None, None, None, None, "", "", ""),
    ]
    for i, row in enumerate(cavs, start=5):
        for j, val in enumerate(row, start=1):
            cell = ws4.cell(row=i, column=j, value=val if val is not None else "")
            if j == 1:
                cell.alignment = left
            else:
                cell.alignment = center
    apply_body_border(ws4, 5, 4 + len(cavs), 1, len(headers4))
    widths4 = [24, 8, 12, 10, 10, 8, 10, 8, 8, 8]
    for i, w in enumerate(widths4, start=1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4["A10"] = "※ミッチェルの6スティールはキャリア最多（プレーオフ）。空欄は確認できなかった項目（推測しない方針）。"
    ws4["A10"].font = Font(italic=True, size=9, color="666666")
    ws4.merge_cells("A10:J10")

    # Sheet 5: 試合経過のキーポイント
    ws5 = wb.create_sheet("Game1経過")
    ws5["A1"] = "Game 1 試合経過の主要ポイント"
    ws5["A1"].font = Font(bold=True, size=14)
    ws5.merge_cells("A1:B1")

    ws5["A3"] = "時点／項目"
    ws5["B3"] = "内容"
    apply_header(ws5, 3, 2)
    progress = [
        ("前半終了時", "キャバリアーズ 48 - ニックス 46（キャバリアーズが前半終盤に18-6のラン）"),
        ("第4Q 残り7:52", "キャバリアーズ 93 - ニックス 71（キャバリアーズ22点リード）"),
        ("第4Q残り7:52～OT終了", "ニックスが44-11のランで逆転勝利"),
        ("OT 14得点中", "OG・アヌノビーが9点を獲得"),
        ("歴史的記録", "プレーバイプレー記録が残る1997年以降、カンファレンスファイナルにおける第4Q最大の逆転（NBA公式）"),
        ("第4Q以降のキャバリアーズFG成功率", "残り時間で約22%"),
    ]
    for i, (k, v) in enumerate(progress, start=4):
        ws5.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws5.cell(row=i, column=2, value=v)
    apply_body_border(ws5, 4, 3 + len(progress), 1, 2)
    ws5.column_dimensions["A"].width = 30
    ws5.column_dimensions["B"].width = 70

    # Sheet 6: インジュリーレポート
    ws6 = wb.create_sheet("インジュリー")
    ws6["A1"] = "Game 2 インジュリーレポート"
    ws6["A1"].font = Font(bold=True, size=14)
    ws6.merge_cells("A1:C1")
    headers6 = ["選手", "チーム", "状況"]
    for i, h in enumerate(headers6, start=1):
        ws6.cell(row=3, column=i, value=h)
    apply_header(ws6, 3, 3)
    injury = [
        ("OG・アヌノビー", "ニックス", "出場可能。ただし5月6日のハムストリング肉離れから完全回復には至っていない"),
        ("ラリー・ナンス・ジュニア", "キャバリアーズ", "questionable（出場微妙）"),
    ]
    for i, row in enumerate(injury, start=4):
        for j, val in enumerate(row, start=1):
            cell = ws6.cell(row=i, column=j, value=val)
            cell.alignment = left
    apply_body_border(ws6, 4, 3 + len(injury), 1, 3)
    ws6.column_dimensions["A"].width = 26
    ws6.column_dimensions["B"].width = 18
    ws6.column_dimensions["C"].width = 60
    ws6["A7"] = "※それ以外の両チームの主要選手は出場可能。"
    ws6["A7"].font = Font(italic=True, size=9, color="666666")
    ws6.merge_cells("A7:C7")

    # Sheet 7: 出典
    ws7 = wb.create_sheet("出典")
    ws7["A1"] = "出典・参考リンク"
    ws7["A1"].font = Font(bold=True, size=14)
    ws7.merge_cells("A1:B1")
    ws7["A3"] = "媒体"
    ws7["B3"] = "URL"
    apply_header(ws7, 3, 2)
    sources = [
        ("NBA.com 2026 East Final 公式", "https://www.nba.com/playoffs/2026/east-final"),
        ("NBA.com ニックス公式 Game 2 ゲームノート", "https://www.nba.com/knicks/news/knicks-vs-cavaliers-game-notes-may-21-2026-eastern-conference-finals-game-2"),
        ("NBA.com Game 1 ライブアップデート", "https://www.nba.com/news/live-updates-2026-nba-playoffs-eastern-conference-finals-cavaliers-knicks-in-the-garden"),
        ("NBA.com 4 takeaways Game 1", "https://www.nba.com/news/4-takeaways-cavaliers-knicks-game-1"),
        ("ESPN Game 1 ファイナルスコア", "https://www.espn.com/nba/game/_/gameId/401873341/cavaliers-knicks"),
        ("ESPN 2026 NBAプレーオフ スケジュール", "https://www.espn.com/nba/story/_/id/48419498/nba-playoffs-2026-play-finals-schedule-scores-news-highlights-bracket-dates"),
        ("Yahoo Sports Game 1 リキャップ", "https://sports.yahoo.com/articles/knicks-vs-cavaliers-live-score-235147883.html"),
        ("amNewYork Game 1 リキャップ", "https://www.amny.com/sports/knicks-cavaliers-game1-ecf-5-19-26/"),
    ]
    for i, (k, v) in enumerate(sources, start=4):
        ws7.cell(row=i, column=1, value=k)
        ws7.cell(row=i, column=2, value=v)
    apply_body_border(ws7, 4, 3 + len(sources), 1, 2)
    ws7.column_dimensions["A"].width = 40
    ws7.column_dimensions["B"].width = 95

    path = f"{OUT_DIR}/NBA_ECF_Game2_Knicks_vs_Cavaliers.xlsx"
    wb.save(path)
    return path


if __name__ == "__main__":
    w = make_word()
    x = make_excel()
    print("Word:", w)
    print("Excel:", x)
